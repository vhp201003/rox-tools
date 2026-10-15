import sqlite3
import openpyxl
import re
import sys
from pathlib import Path
import UnityPy

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / 'data' / 'game_data.db'
EXCEL_PATH = ROOT / 'assets' / 'templates' / 'arang-simulator-master.xlsx'
STREAMING_ASSETS = ROOT / 'references' / 'RO_SEA' / 'RO_SEAGame' / 'rox_Data' / 'StreamingAssets'

def find_bundle(name):
    bl_path = STREAMING_ASSETS / 'Windows' / 'BundleList.txt'
    with open(bl_path, 'r', encoding='utf-8') as f:
        for line in f:
            parts = line.strip().split(',')
            if parts[0] == name:
                return bl_path.parent / f'{parts[1]}.bundle'
    return None

def get_text(bpath, aname):
    env = UnityPy.load(str(bpath))
    for obj in env.objects:
        if obj.type.name == 'TextAsset':
            d = obj.read()
            if getattr(d, 'm_Name', '') == aname:
                s = getattr(d, 'm_Script', b'')
                return s.decode('utf-8', errors='ignore') if isinstance(s, bytes) else str(s)
    return ''

def extract_excel_val(cell):
    v = cell.value
    if v is None:
        return None
    s = str(v)
    if 'COMPUTED_VALUE' in s or 'IMPORTRANGE' in s:
        m = re.search(r',\"?(.*?)\"?\)$', s)
        if m:
            val = m.group(1).rstrip('"')
            try:
                if '.' in val or 'E' in val or 'e' in val:
                    return float(val)
                return int(val)
            except ValueError:
                return val
    return v

def norm_card(name):
    n = name.lower()
    n = re.sub(r'\[.*?\]', '', n)
    n = re.sub(r'\(.*?\)', '', n)
    n = n.replace('card', '')
    n = re.sub(r'[^\w\s]', '', n)
    return ' '.join(n.split())

def enrich_cards():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute('PRAGMA journal_mode = WAL;')

    bpath0 = find_bundle('lua_data_data_0.bundle')
    bpath1 = find_bundle('lua_data_data_1.bundle')
    en_path = find_bundle('lua_data_en.bundle')

    propid_str = get_text(bpath0, 'data_PropID')
    cardadv_str = get_text(bpath1, 'data_ItemV2_CardAdvAttr')
    en_str = get_text(en_path, 'en_langs')

    prop_meta = {}
    blocks = re.findall(r'\[(\d+)\]=\{([^}]+)\}', propid_str)
    for pid_str, body in blocks:
        pid = int(pid_str)
        if_pct = 1 if '["ifPercentage"]=1' in body else 0
        m_name = re.search(rf'\["Property{pid}"\]\s*=\s*"([^"]*)"', en_str)
        eng_name = m_name.group(1) if m_name else f"Prop_{pid}"
        prop_meta[pid] = {"name": eng_name, "if_pct": if_pct}

    card_adv_map = {}
    adv_blocks = re.findall(r'\[(\d+)\]\s*=\s*\{([\s\S]*?\["staticId"\]\s*=\s*\1[\s\S]*?)\n\s*\},', cardadv_str)
    for cid_str, body in adv_blocks:
        cid = int(cid_str)
        m_ids = re.search(r'\["attrId"\]\s*=\s*\{\s*"([^"{}]+)"', body)
        m_vals = re.search(r'\["attrValue"\]\s*=\s*\{\s*"([^"{}]+)"', body)
        if m_ids and m_vals:
            ids = [int(x) for x in m_ids.group(1).split('|') if x.strip()]
            vals = [float(x) for x in m_vals.group(1).split('|') if x.strip()]
            stats = []
            for pid, v in zip(ids, vals):
                meta = prop_meta.get(pid, {"name": f"Prop_{pid}", "if_pct": 0})
                pname = meta["name"]
                if meta["if_pct"] == 1 or (0 < abs(v) < 1):
                    val_str = f"+{v*100:.1f}%".replace(".0%", "%")
                else:
                    val_str = f"+{int(v) if v.is_integer() else v}"
                stats.append(f"{pname} {val_str}")
            card_adv_map[cid] = " | ".join(stats)

    cardskill_map = {}
    for m in re.finditer(r'\["cardskill_(\d+)"\]\s*=\s*"([^"]*)"', en_str):
        cardskill_map[int(m.group(1))] = m.group(2).replace(r'\n', ' ').strip()

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
    ws = wb['MasterData']

    slots = [
        ('Weapon', 24, 25, 29),
        ('Clothes', 34, 35, 39),
        ('Cloak', 44, 45, 49),
        ('Shoes', 54, 55, 59),
        ('Accessory', 64, 65, 69),
        ('Headgear', 74, 75, 79),
    ]

    master_cards = {}
    for slot_name, name_col, sname_start, sval_start in slots:
        for r in range(4, ws.max_row + 1):
            cname = extract_excel_val(ws.cell(row=r, column=name_col))
            if not cname or str(cname).strip() == '' or str(cname).startswith('None'):
                continue
            cname_clean = str(cname).strip()
            stats = []
            for i in range(4):
                sname = extract_excel_val(ws.cell(row=r, column=sname_start + i))
                sval = extract_excel_val(ws.cell(row=r, column=sval_start + i))
                if sname and sval is not None:
                    if isinstance(sval, float) and 0 < sval < 1:
                        val_str = f'+{sval*100:.1f}%'.replace('.0%', '%')
                    elif isinstance(sval, (int, float)):
                        val_str = f'+{sval}'
                    else:
                        val_str = str(sval)
                    stats.append(f'{sname} {val_str}')
            stat_desc = ' | '.join(stats) if stats else ''
            master_cards[cname_clean.lower()] = stat_desc

    master_lookup = {}
    for k, v in master_cards.items():
        if v:
            master_lookup[k] = v
            master_lookup[k.replace(' card', '')] = v
            master_lookup[norm_card(k)] = v

    special_passives = {
        10101057: "Normal attacks deal Splash Physical Damage to nearby enemies within range",
        10101043: "Deals 100% damage to Large, Medium, and Small monsters",
        10104015: "Magic Damage Reflect +10% | M.DEF +10",
        10105015: "Variable Cast Time -10% | Max HP -10% | Max SP -10%",
        10105016: "Gain the ability to detect Hidden/Stealthed units",
        10106031: "When dealing normal attacks, there is a chance to cast Lv.5 Lex Aeterna",
        10107019: "When taking physical damage, there is a chance to inflict Petrification on enemies",
        10107020: "P.ATK +20 | Increase Physical Damage to Demi-Human monsters +10%",
        10107021: "Physical Damage to Medium-sized monsters +10%",
    }

    aliases = {
        "falling bishop": "fallen bishop",
        "cat o' nine tail": "cat o' nine tails",
        "maya puple": "maya purple",
        "ogretooth": "ogre tooth",
        "ancient mummy": "mummy",
    }

    cur.execute('SELECT id, name, attr_desc FROM cards')
    rows = cur.fetchall()

    updates = []
    before_count = sum(1 for r in rows if r[2] and r[2].strip())

    for cid, cname, current_attr in rows:
        final_attr = current_attr.strip() if current_attr else ""

        if not final_attr and cid in card_adv_map and card_adv_map[cid]:
            final_attr = card_adv_map[cid]

        if not final_attr and cid in cardskill_map and cardskill_map[cid]:
            final_attr = cardskill_map[cid]

        if not final_attr and cid in special_passives:
            final_attr = special_passives[cid]

        if not final_attr:
            clean_name = cname.strip().lower()
            for a_from, a_to in aliases.items():
                if a_from in clean_name:
                    clean_name = clean_name.replace(a_from, a_to)
                    break
            k_exact = clean_name
            k_no_card = k_exact.replace(' card', '')
            k_norm = norm_card(clean_name)
            if k_exact in master_lookup:
                final_attr = master_lookup[k_exact]
            elif k_no_card in master_lookup:
                final_attr = master_lookup[k_no_card]
            elif k_norm in master_lookup:
                final_attr = master_lookup[k_norm]

        if not final_attr and cid > 1000000000:
            base_id = cid // 100
            if base_id in card_adv_map and card_adv_map[base_id]:
                final_attr = f"{card_adv_map[base_id]} [★ Awakened]"
            elif base_id in special_passives:
                final_attr = f"{special_passives[base_id]} [★ Awakened]"

        if final_attr and final_attr != current_attr:
            updates.append((final_attr, cid))

    cur.executemany('UPDATE cards SET attr_desc = ? WHERE id = ?', updates)
    conn.commit()

    cur.execute('SELECT count(*) FROM cards WHERE attr_desc IS NOT NULL AND trim(attr_desc) != ""')
    after_count = cur.fetchone()[0]

    print(f"Updated {len(updates)} cards.")
    print(f"Before: {before_count}/{len(rows)} cards ({before_count/len(rows)*100:.1f}%)")
    print(f"After: {after_count}/{len(rows)} cards ({after_count/len(rows)*100:.1f}%)")

    conn.close()

if __name__ == '__main__':
    enrich_cards()
