import json
import re
import sqlite3
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / "data" / "game_data.db"
EXCEL_PATH = ROOT / "assets" / "templates" / "arang-simulator-master.xlsx"

SLOT_NAMES = {
    1: "Weapon",
    2: "Armor",
    3: "Cloak",
    4: "Shoes",
    5: "Headgear",
    6: "Accessory"
}

def extract_val(cell):
    v = cell.value
    if v is None:
        return None
    s = str(v)
    if 'COMPUTED_VALUE' in s or 'IMPORTRANGE' in s:
        m = re.search(r',\"?(.*?)\"?\)$', s)
        if m:
            val = m.group(1).rstrip('"')
            try:
                if '.' in val or 'E' in val or 'e' in val:
                    return float(val)
                return int(val)
            except ValueError:
                return val
    return v

def migrate():
    print(f"Connecting to {DB_PATH}...")
    conn = sqlite3.connect(str(DB_PATH))
    conn.execute("PRAGMA journal_mode=WAL;")
    conn.execute("PRAGMA synchronous=NORMAL;")
    cur = conn.cursor()

    # 1. Update monsters schema
    cur.execute("PRAGMA table_xinfo(monsters)")
    m_cols = [r[1] for r in cur.fetchall()]
    
    base_ext_cols = [
        "def_base", "def_ext", 
        "mdef_base", "mdef_ext", 
        "crit_res_base", "crit_res_ext", 
        "flee_base", "flee_ext"
    ]
    for col in base_ext_cols:
        if col not in m_cols:
            print(f"Adding column {col} to monsters...")
            cur.execute(f"ALTER TABLE monsters ADD COLUMN {col} REAL DEFAULT 0.0")

    conn.commit()

    cur.execute("PRAGMA table_xinfo(monsters)")
    m_cols = [r[1] for r in cur.fetchall()]

    gen_cols = [
        ("total_def", "ROUND(def_base + def_ext, 2)"),
        ("total_mdef", "ROUND(mdef_base + mdef_ext, 2)"),
        ("total_crit_res", "ROUND(crit_res_base + crit_res_ext, 2)"),
        ("total_flee", "ROUND(flee_base + flee_ext, 2)")
    ]
    for col, expr in gen_cols:
        if col not in m_cols:
            print(f"Adding generated column {col} to monsters...")
            cur.execute(f"ALTER TABLE monsters ADD COLUMN {col} REAL GENERATED ALWAYS AS ({expr})")

    conn.commit()

    # 2. Extract stats from Excel Master sheet MonsterData
    print("Loading Arang Simulator Master sheet MonsterData...")
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=False)
    ws = wb['MonsterData']

    excel_records = []
    for r in range(4, ws.max_row + 1):
        name = extract_val(ws.cell(row=r, column=2))
        if not name or str(name).startswith("Average") or str(name).startswith("Last updated"):
            continue
        level = extract_val(ws.cell(row=r, column=1))
        hp = extract_val(ws.cell(row=r, column=12))
        
        def_base = extract_val(ws.cell(row=r, column=20)) or 0.0
        def_ext = extract_val(ws.cell(row=r, column=21)) or 0.0
        mdef_base = extract_val(ws.cell(row=r, column=22)) or 0.0
        mdef_ext = extract_val(ws.cell(row=r, column=23)) or 0.0
        crit_res_base = extract_val(ws.cell(row=r, column=24)) or 0.0
        crit_res_ext = extract_val(ws.cell(row=r, column=25)) or 0.0
        flee_base = extract_val(ws.cell(row=r, column=26)) or 0.0
        flee_ext = extract_val(ws.cell(row=r, column=27)) or 0.0

        excel_records.append({
            "name": str(name).strip(),
            "level": int(float(level)) if level is not None else 0,
            "hp": int(float(hp)) if hp is not None else 0,
            "def_base": round(float(def_base) * 100, 2),
            "def_ext": round(float(def_ext) * 100, 2),
            "mdef_base": round(float(mdef_base) * 100, 2),
            "mdef_ext": round(float(mdef_ext) * 100, 2),
            "crit_res_base": round(float(crit_res_base) * 100, 2),
            "crit_res_ext": round(float(crit_res_ext) * 100, 2),
            "flee_base": round(float(flee_base) * 100, 2),
            "flee_ext": round(float(flee_ext) * 100, 2),
        })

    print(f"Extracted {len(excel_records)} records from Excel. Updating monsters table...")

    stats_updates = 0
    for rec in excel_records:
        # Match by name and level first
        cur.execute("SELECT id FROM monsters WHERE LOWER(name) = LOWER(?) AND level = ?", (rec["name"], rec["level"]))
        matched_ids = [r[0] for r in cur.fetchall()]
        if not matched_ids:
            cur.execute("SELECT id FROM monsters WHERE LOWER(name) = LOWER(?)", (rec["name"],))
            matched_ids = [r[0] for r in cur.fetchall()]
        
        for mid in matched_ids:
            cur.execute("""
                UPDATE monsters 
                SET def_base = ?, def_ext = ?,
                    mdef_base = ?, mdef_ext = ?,
                    crit_res_base = ?, crit_res_ext = ?,
                    flee_base = ?, flee_ext = ?,
                    max_hp = CASE WHEN ? > 0 THEN ? ELSE max_hp END
                WHERE id = ?
            """, (
                rec["def_base"], rec["def_ext"],
                rec["mdef_base"], rec["mdef_ext"],
                rec["crit_res_base"], rec["crit_res_ext"],
                rec["flee_base"], rec["flee_ext"],
                rec["hp"], rec["hp"],
                mid
            ))
            stats_updates += 1

    conn.commit()
    print(f"Updated combat stats for {stats_updates} monster instances.")

    # 3. Create monster_cards table
    print("Creating monster_cards table...")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS monster_cards (
            monster_id INTEGER NOT NULL,
            card_id INTEGER NOT NULL,
            weather_phase TEXT NOT NULL DEFAULT 'Normal' CHECK(weather_phase IN ('Normal', 'Day', 'Night', 'Void')),
            slot_id INTEGER DEFAULT 0,
            slot_name TEXT DEFAULT '',
            PRIMARY KEY (monster_id, card_id),
            FOREIGN KEY (monster_id) REFERENCES monsters(id) ON DELETE CASCADE,
            FOREIGN KEY (card_id) REFERENCES cards(id) ON DELETE CASCADE
        );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mc_monster ON monster_cards(monster_id);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mc_phase ON monster_cards(weather_phase);")
    cur.execute("DELETE FROM monster_cards;")
    conn.commit()

    # 4. Populate weather cards (Day, Night, Void)
    cur.execute("SELECT id, name, card_type, monster_id FROM cards WHERE name LIKE '%[Day]%' OR name LIKE '%[Night]%' OR name LIKE '%[Void]%'")
    weather_cards = cur.fetchall()
    
    weather_links = []
    for cid, name, card_type, mid in weather_cards:
        phase_m = re.search(r'\[(Day|Night|Void)\]', name, re.IGNORECASE)
        phase = phase_m.group(1).capitalize() if phase_m else 'Normal'
        slot_name = SLOT_NAMES.get(card_type, "Unknown")
        base_name = re.sub(r'\s*Card\s*\[(Day|Night|Void)\]', '', name, flags=re.IGNORECASE).strip()

        # Find target monster id
        cur.execute("SELECT id, is_mvp, is_mini FROM monsters WHERE id = ?", (mid,))
        m_row = cur.fetchone()
        
        target_ids = []
        if m_row and (m_row[1] or m_row[2] or 10001 <= m_row[0] <= 39999):
            target_ids.append(m_row[0])
        else:
            cur.execute("SELECT id FROM monsters WHERE LOWER(name) = LOWER(?) ORDER BY is_mvp DESC, is_mini DESC, id ASC", (base_name,))
            rows = cur.fetchall()
            if rows:
                target_ids.append(rows[0][0])
            else:
                cur.execute("SELECT id FROM monsters WHERE LOWER(name) LIKE LOWER(?) ORDER BY is_mvp DESC, is_mini DESC, id ASC", (f"{base_name}%",))
                frows = cur.fetchall()
                if frows:
                    target_ids.append(frows[0][0])

        for tid in target_ids:
            weather_links.append((tid, cid, phase, card_type, slot_name))

    cur.executemany("""
        INSERT OR IGNORE INTO monster_cards (monster_id, card_id, weather_phase, slot_id, slot_name)
        VALUES (?, ?, ?, ?, ?)
    """, weather_links)
    conn.commit()
    print(f"Inserted {len(weather_links)} weather card links.")

    # 5. Populate Normal cards (for normal monsters)
    cur.execute("""
        SELECT id, name, card_type, monster_id 
        FROM cards 
        WHERE name NOT LIKE '%[Day]%' 
          AND name NOT LIKE '%[Night]%' 
          AND name NOT LIKE '%[Void]%'
          AND name NOT LIKE 'Card #%'
    """)
    normal_cards = cur.fetchall()

    normal_links = []
    for cid, name, card_type, mid in normal_cards:
        slot_name = SLOT_NAMES.get(card_type, "Unknown")
        base_name = re.sub(r'\s*Card$', '', name, flags=re.IGNORECASE).strip()

        target_ids = []
        if mid and mid > 0:
            cur.execute("SELECT id FROM monsters WHERE id = ?", (mid,))
            if cur.fetchone():
                target_ids.append(mid)
        
        if not target_ids:
            cur.execute("SELECT id FROM monsters WHERE LOWER(name) = LOWER(?) ORDER BY is_canonical DESC, id ASC", (base_name,))
            rows = cur.fetchall()
            if rows:
                target_ids.append(rows[0][0])

        for tid in target_ids:
            normal_links.append((tid, cid, 'Normal', card_type, slot_name))

    cur.executemany("""
        INSERT OR IGNORE INTO monster_cards (monster_id, card_id, weather_phase, slot_id, slot_name)
        VALUES (?, ?, ?, ?, ?)
    """, normal_links)
    conn.commit()
    print(f"Inserted {len(normal_links)} normal card links.")

    # 6. Verify Agav Lv 150 benchmark
    print("\n--- Verifying Benchmark: Agav Lv 150 ---")
    cur.execute("""
        SELECT id, name, level, race_name, element_name, max_hp, 
               def_base, def_ext, total_def,
               mdef_base, mdef_ext, total_mdef,
               crit_res_base, crit_res_ext, total_crit_res,
               flee_base, flee_ext, total_flee
        FROM monsters 
        WHERE LOWER(name) = 'agav'
    """)
    agav_rows = cur.fetchall()
    for r in agav_rows:
        print(f"Monster ID: {r[0]} | {r[1]} Lv.{r[2]} ({r[3]} / {r[4]})")
        print(f"  Health: {r[5]:,}")
        print(f"  Total Def: {r[8]:.2f}% ({r[6]:.2f}% + {r[7]:.2f}%)")
        print(f"  Total M.Def: {r[11]:.2f}% ({r[9]:.2f}% + {r[10]:.2f}%)")
        print(f"  Total Crit.Res: {r[14]:.2f}% ({r[12]:.2f}% + {r[13]:.2f}%)")
        print(f"  Total Flee: {r[17]:.2f}% ({r[15]:.2f}% + {r[16]:.2f}%)")

    # 7. Verify Angeling and Armored Titan cards in monster_cards
    print("\n--- Verifying Angeling Cards in monster_cards ---")
    cur.execute("""
        SELECT mc.monster_id, m.name, c.id, c.name, mc.weather_phase, mc.slot_name, c.attr_desc
        FROM monster_cards mc
        JOIN monsters m ON m.id = mc.monster_id
        JOIN cards c ON c.id = mc.card_id
        WHERE m.name = 'Angeling'
        ORDER BY mc.weather_phase
    """)
    for r in cur.fetchall():
        print(f"  Boss: {r[1]} -> [{r[4]}] {r[3]} | Slot: {r[5]} | FX: {r[6]}")

    print("\n--- Verifying Armored Titan Cards in monster_cards ---")
    cur.execute("""
        SELECT mc.monster_id, m.name, c.id, c.name, mc.weather_phase, mc.slot_name, c.attr_desc
        FROM monster_cards mc
        JOIN monsters m ON m.id = mc.monster_id
        JOIN cards c ON c.id = mc.card_id
        WHERE m.name = 'Armored Titan'
        ORDER BY mc.weather_phase
    """)
    for r in cur.fetchall():
        print(f"  Boss: {r[1]} -> [{r[4]}] {r[3]} | Slot: {r[5]} | FX: {r[6]}")

    conn.close()
    print("\nMigration completed successfully!")

if __name__ == "__main__":
    migrate()
