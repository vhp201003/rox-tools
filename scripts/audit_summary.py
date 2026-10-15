import openpyxl

wb = openpyxl.load_workbook('assets/templates/arang-simulator-master.xlsx', data_only=False)

def summarize_sheet(sheet_name):
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]
    formula_count = 0
    val_count = 0
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if cell.value is not None:
                val_str = str(cell.value)
                if val_str.startswith('='):
                    formula_count += 1
                else:
                    val_count += 1
    print(f"Sheet '{sheet_name}': {ws.max_row} rows x {ws.max_column} cols | Values: {val_count}, Formulas: {formula_count}")

print("=== WORKBOOK SHEET SUMMARY ===")
for name in wb.sheetnames:
    summarize_sheet(name)
